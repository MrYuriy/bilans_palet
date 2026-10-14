"""XLSX files: the daily balance (the D-1 report for LM) and the movement register.

Dates are real Excel dates and quantities real numbers, so the file can be
filtered and summed as it is.
"""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.core.timeutils import local_now
from app.schemas.balance import DailyBalanceOut, DayValues
from app.schemas.movement import MovementOut

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

_BOLD = Font(bold=True)
_HEADER_FILL = PatternFill("solid", fgColor="EEF0EC")
_TOTAL_FILL = PatternFill("solid", fgColor="EEF5F1")
_DATE = "DD.MM.YYYY"
_DATETIME = "DD.MM.YYYY HH:MM"


def _header(ws: Worksheet, row: int, titles: list[str]) -> None:
    for col, title in enumerate(titles, start=1):
        cell = ws.cell(row=row, column=col, value=title)
        cell.font = _BOLD
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def _widths(ws: Worksheet, widths: list[int]) -> None:
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width


def _save(wb: Workbook) -> bytes:
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# ---------- daily balance ----------

DAY_COLUMNS = [
    "Grupa",
    "Typ nośnika",
    "Stan otwarcia",
    "Przyjęcia",
    "Wydania",
    "Korekty",
    "Przeksięgowania",
    "Stan zamknięcia",
]


def _day_values(r: DayValues) -> list[int]:
    return [
        r.opening,
        r.receipts,
        r.shipments,
        r.corrections,
        r.reclassifications,
        r.closing,
    ]


def daily_balance_xlsx(data: DailyBalanceOut) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = f"Bilans {data.day.strftime('%d.%m.%Y')}"

    status = (
        f"zamknięta {data.closed_at.strftime('%d.%m.%Y %H:%M')}"
        if data.closed
        else "podgląd — doba otwarta, liczby mogą się zmienić"
    )
    ws["A1"] = "Bilans dzienny nośników"
    ws["A1"].font = Font(bold=True, size=14)
    info = [
        ("Doba", data.day),
        ("Status", status),
        ("Wygenerowano", local_now().replace(microsecond=0)),
    ]
    for offset, (label, value) in enumerate(info, start=2):
        ws.cell(row=offset, column=1, value=label).font = _BOLD
        cell = ws.cell(row=offset, column=2, value=value)
        if offset == 2:
            cell.number_format = _DATE
        if label == "Wygenerowano":
            cell.number_format = _DATETIME

    ws.cell(
        row=6,
        column=1,
        value="Stan zamknięcia = otwarcie + przyjęcia − wydania + korekty "
        "+ przeksięgowania. Korekty obejmują inwentaryzację. "
        "Razem grupy liczy tylko typy liczone w raportach.",
    )

    header_row = 8
    _header(ws, header_row, DAY_COLUMNS)
    row = header_row + 1
    for group in data.groups:
        for r in group.rows:
            ws.append([group.group.name, r.carrier_type.name, *_day_values(r)])
            row += 1
        ws.append([group.group.name, "RAZEM", *_day_values(group.total)])
        for col in range(1, len(DAY_COLUMNS) + 1):
            ws.cell(row=row, column=col).font = _BOLD
            ws.cell(row=row, column=col).fill = _TOTAL_FILL
        row += 1

    _widths(ws, [16, 20, 14, 12, 12, 12, 16, 16])
    return _save(wb)


# ---------- movement register ----------

MOVEMENT_COLUMNS = [
    "Data operacji",
    "Operacja",
    "Typ nośnika",
    "Zmiana stanu",
    "Nr dokumentu",
    "Komentarz",
    "Użytkownik",
    "Wprowadzono",
]


def movements_xlsx(movements: list[MovementOut]) -> bytes:
    """One row per movement line; a stock count that matched has one row with 0."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Rejestr operacji"
    _header(ws, 1, MOVEMENT_COLUMNS)

    row = 2
    for m in movements:
        lines = m.lines or [None]
        for line in lines:
            ws.append(
                [
                    m.operation_date,
                    m.type_label,
                    line.carrier_type.name if line else "—",
                    line.quantity if line else 0,
                    m.document_number or "",
                    m.comment or "",
                    m.user.username,
                    m.created_at.replace(microsecond=0),
                ]
            )
            ws.cell(row=row, column=1).number_format = _DATE
            ws.cell(row=row, column=8).number_format = _DATETIME
            row += 1

    _widths(ws, [13, 28, 20, 13, 18, 30, 16, 17])
    ws.auto_filter.ref = (
        f"A1:{get_column_letter(len(MOVEMENT_COLUMNS))}{max(row - 1, 1)}"
    )
    return _save(wb)
