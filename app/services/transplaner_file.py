"""Reading the monthly carrier report from Transplaner (raport rozliczeniowy
nośników).

The file holds one month, day by day, per recipient (store, hub, ...). Its
sheets are found by cell A1, e.g. "PAL EUR/ EPAL - wyjścia 2026-10" (shipments
to the stores) and "PAL EUR/ EPAL - wejścia 2026-10" (returns); a header row
starting with "Sklep" names the day columns 1..N and the "Suma" column.
Anything after "Suma" (e.g. the returns by receipt date) is not read.
"""

import calendar
import re
from dataclasses import dataclass, field
from datetime import date
from io import BytesIO
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException

from app.core.exc import BadRequestException

SHIPMENTS = "wyjścia"
RETURNS = "wejścia"

_TITLE = re.compile(
    r"^\s*(?P<carrier>.+?)\s+-\s+(?P<kind>wyjścia|wejścia)\s+(?P<year>\d{4})-(?P<month>\d{2})\s*$"
)
# The header row is looked for among the first rows only.
_HEADER_SEARCH_ROWS = 10


@dataclass
class Sheet:
    title: str
    carrier: str
    kind: str  # SHIPMENTS or RETURNS
    month: date  # its first day
    # recipient -> day -> quantity (only days with a quantity)
    rows: dict[str, dict[date, int]] = field(default_factory=dict)


@dataclass
class TransplanerFile:
    carrier: str
    month: date
    sheets: dict[str, Sheet]  # kind -> sheet


def _fail(message: str):
    raise BadRequestException(message)


def _quantity(value, where: str) -> int:
    """A cell of the day grid: empty = 0, else a whole number >= 0."""
    if value is None or value == "":
        return 0
    if isinstance(value, bool):
        _fail(f"{where}: „{value}” nie jest liczbą")
    if isinstance(value, int | float) and value == int(value) and value >= 0:
        return int(value)
    _fail(f"{where}: „{value}” nie jest liczbą całkowitą ≥ 0")


def _read_sheet(ws, title: str, carrier: str, kind: str, month: date) -> Sheet:
    rows = list(ws.iter_rows(values_only=True))
    header_index = next(
        (
            i
            for i, row in enumerate(rows[:_HEADER_SEARCH_ROWS])
            if row and str(row[0] or "").strip() == "Sklep"
        ),
        None,
    )
    if header_index is None:
        _fail(f"Zakładka „{title}”: brak wiersza nagłówka „Sklep”")

    header = rows[header_index]
    days_in_month = calendar.monthrange(month.year, month.month)[1]
    day_columns: list[int] = []
    sum_column = None
    for col, value in enumerate(header[1:], start=1):
        if isinstance(value, str) and value.strip() == "Suma":
            sum_column = col
            break
        if value != len(day_columns) + 1:
            _fail(
                f"Zakładka „{title}”: kolumna {get_column_letter(col + 1)} nagłówka"
                f" powinna mieć numer dnia {len(day_columns) + 1}"
            )
        day_columns.append(col)
    if sum_column is None or not day_columns:
        _fail(f"Zakładka „{title}”: brak kolumn dni i kolumny „Suma”")

    sheet = Sheet(title=title, carrier=carrier, kind=kind, month=month)
    for number, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
        name = str(row[0] or "").strip() if row else ""
        if not name:
            continue
        days: dict[date, int] = {}
        for col in day_columns:
            where = f"Zakładka „{title}”, komórka {get_column_letter(col + 1)}{number}"
            quantity = _quantity(row[col] if col < len(row) else None, where)
            if not quantity:
                continue
            if col > days_in_month:
                _fail(f"{where}: dzień {col} nie istnieje w tym miesiącu")
            days[date(month.year, month.month, col)] = quantity

        total = row[sum_column] if sum_column < len(row) else None
        if total is not None:
            where = f"Zakładka „{title}”, wiersz {number} ({name})"
            expected = _quantity(total, where)
            if expected != sum(days.values()):
                _fail(
                    f"{where}: suma dni {sum(days.values())} nie zgadza się"
                    f" z kolumną „Suma” ({expected})"
                )
        if days:
            merged = sheet.rows.setdefault(name, {})
            for day, quantity in days.items():
                merged[day] = merged.get(day, 0) + quantity
    return sheet


def read_transplaner(content: bytes) -> TransplanerFile:
    """Read the shipments and returns sheets; refuses a file that does not add up."""
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException, KeyError, ValueError, OSError):
        _fail("To nie jest plik Excel (.xlsx)")

    sheets: dict[str, Sheet] = {}
    try:
        for ws in workbook.worksheets:
            first = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
            match = _TITLE.match(str(first[0] or "")) if first else None
            if match is None:
                continue
            kind = match["kind"]
            if kind in sheets:
                _fail(f"Plik ma dwie zakładki „{kind}”")
            month = date(int(match["year"]), int(match["month"]), 1)
            sheets[kind] = _read_sheet(
                ws, ws.title, match["carrier"].strip(), kind, month
            )
    finally:
        workbook.close()

    if not sheets:
        _fail(
            "W pliku nie ma zakładki wyjść ani zwrotów"
            " (w polu A1 np. „PAL EUR/ EPAL - wyjścia 2026-10”)"
        )
    if len({(s.carrier, s.month) for s in sheets.values()}) > 1:
        _fail("Zakładki wyjść i zwrotów dotyczą różnych nośników lub miesięcy")
    any_sheet = next(iter(sheets.values()))
    return TransplanerFile(any_sheet.carrier, any_sheet.month, sheets)
