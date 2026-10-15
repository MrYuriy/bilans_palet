"""The Transplaner file: reading it, previewing it, booking only what changed."""

from datetime import date, timedelta
from io import BytesIO

import pytest
from openpyxl import Workbook

from app.core.exc import BadRequestException
from app.services.transplaner_file import read_transplaner
from tests.conftest import login, make_user, stock_of, warehouse_today

TODAY = warehouse_today()
# A month that has fully passed, so its days never depend on today.
LAST_MONTH = (TODAY.replace(day=1) - timedelta(days=1)).replace(day=1)
UKR = "Leroy Merlin Ukraine LLC"


def day(n: int, month: date = LAST_MONTH) -> date:
    return month.replace(day=n)


def build(
    shipments: dict | None = None,
    returns: dict | None = None,
    month: str | None = None,
    carrier: str = "PAL EUR/ EPAL",
    returns_month: str | None = None,
    tamper=None,
) -> bytes:
    """A file shaped like the real report. shipments/returns: recipient ->
    {day number: quantity}."""
    month = month or LAST_MONTH.strftime("%Y-%m")
    wb = Workbook()
    saldo = wb.active
    saldo.title = "PAL EUR- EPAL - saldo"
    saldo["A1"] = f"{carrier} - saldo {month}"

    def sheet(kind, rows, sheet_month):
        ws = wb.create_sheet(f"PAL EUR- EPAL - {kind}")
        ws["A1"] = f"{carrier} - {kind} {sheet_month}"
        if kind == "wejścia":
            ws.append([None, "W/g daty wysłania"])
            ws["AH2"] = "W/g daty przyjęcia"
        header = ["Sklep", *range(1, 32), "Suma"]
        if kind == "wejścia":
            header += [1, 2, 3, 4]
        ws.append(header)
        for name, days in rows.items():
            values = [days.get(n) for n in range(1, 32)]
            row = [name, *values, sum(days.values())]
            if kind == "wejścia":
                row += [999, None, None, None]  # after "Suma": not read
            ws.append(row)
        return ws

    if shipments is not None:
        sheet("wyjścia", shipments, month)
    if returns is not None:
        sheet("wejścia", returns, returns_month or month)
    if tamper:
        tamper(wb)
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# ---------- reading ----------


def test_reading_the_sheets():
    file = read_transplaner(
        build(
            shipments={"002 BIALOLEKA": {1: 26, 2: 15}, UKR: {2: 19}, "TEST": {}},
            returns={"002 BIALOLEKA": {5: 97}},
        )
    )

    assert (file.carrier, file.month) == ("PAL EUR/ EPAL", LAST_MONTH)
    shipments, returns = file.sheets["wyjścia"], file.sheets["wejścia"]
    assert shipments.rows == {
        "002 BIALOLEKA": {day(1): 26, day(2): 15},
        UKR: {day(2): 19},
    }
    # The returns header is one row lower; the block after "Suma" is ignored.
    assert returns.rows == {"002 BIALOLEKA": {day(5): 97}}


@pytest.mark.parametrize(
    ("content", "message"),
    [
        (b"not a spreadsheet", "To nie jest plik Excel (.xlsx)"),
        (
            build(),
            (
                "W pliku nie ma zakładki wyjść ani zwrotów (w polu A1 np."
                " „PAL EUR/ EPAL - wyjścia 2026-10”)"
            ),
        ),
        (
            build(
                shipments={"A": {1: 5}}, returns={"A": {1: 1}}, returns_month="2020-01"
            ),
            "Zakładki wyjść i zwrotów dotyczą różnych nośników lub miesięcy",
        ),
    ],
)
def test_files_that_are_refused(content, message):
    with pytest.raises(BadRequestException) as error:
        read_transplaner(content)
    assert str(error.value) == message


def tampered(cell: str, value):
    def change(wb):
        wb["PAL EUR- EPAL - wyjścia"][cell] = value

    return change


@pytest.mark.parametrize(
    ("cell", "value", "message"),
    [
        ("B3", -4, "komórka B3: „-4” nie jest liczbą całkowitą ≥ 0"),
        ("C3", "dużo", "komórka C3: „dużo” nie jest liczbą całkowitą ≥ 0"),
        ("AG3", 99, "wiersz 3 (A): suma dni 5 nie zgadza się z kolumną „Suma” (99)"),
    ],
)
def test_cells_that_do_not_add_up(cell, value, message):
    content = build(shipments={"A": {1: 5}}, tamper=tampered(cell, value))

    with pytest.raises(BadRequestException) as error:
        read_transplaner(content)
    assert str(error.value) == f"Zakładka „PAL EUR- EPAL - wyjścia”, {message}"


def test_a_day_the_month_does_not_have():
    content = build(shipments={"A": {31: 5}}, month="2026-02")

    with pytest.raises(BadRequestException) as error:
        read_transplaner(content)
    assert str(error.value).endswith("dzień 31 nie istnieje w tym miesiącu")


# ---------- preview and booking ----------


async def preview(client, headers, content):
    return await client.post(
        "/imports/transplaner/preview",
        headers=headers,
        files={"file": ("raport.xlsx", content)},
    )


async def book(client, headers, content, *types):
    types = types or ("STORE_SHIPMENT", "STORE_RETURN", "UKR_SHIPMENT", "UKR_RETURN")
    return await client.post(
        "/imports/transplaner",
        headers=headers,
        files={"file": ("raport.xlsx", content)},
        data={"type": list(types)},
    )


FIRST = build(
    shipments={"002 BIALOLEKA": {1: 26, 2: 15}, "003 GDANSK": {1: 34}, UKR: {2: 19}},
    returns={"002 BIALOLEKA": {1: 10}},
)


async def test_the_preview_shows_what_would_be_booked(client, admin_headers, carriers):
    response = await preview(client, admin_headers, FIRST)

    assert response.status_code == 200, response.text
    body = response.json()
    assert (body["carrier"], body["carrier_type"]["name"]) == ("PAL EUR/ EPAL", "EURO")
    assert body["month"] == LAST_MONTH.isoformat()
    assert body["recipients"] == 3
    assert body["last_day"] == day(2).isoformat()
    kinds = {k["type"]: (k["in_file"], k["booked"], k["change"]) for k in body["kinds"]}
    assert kinds == {
        "STORE_SHIPMENT": (75, 0, -75),
        "UKR_SHIPMENT": (19, 0, -19),
        "STORE_RETURN": (10, 0, 10),
        "UKR_RETURN": (0, 0, 0),
    }
    assert body["warnings"] == [
        f"Plik dotyczy poprzedniego miesiąca ({LAST_MONTH:%m.%Y})"
    ]
    # Nothing is booked by a preview.
    assert (await stock_of(client, admin_headers))["EURO"] == 0


async def test_booking_dates_the_change_today(client, admin_headers, carriers):
    response = await book(client, admin_headers, FIRST)

    assert response.status_code == 200, response.text
    body = response.json()
    assert (body["source"], body["file_name"]) == ("transplaner", "raport.xlsx")
    assert body["summary"]["kinds"]["STORE_SHIPMENT"]["change"] == -75
    movements = (await client.get("/movements", headers=admin_headers)).json()
    booked = {m["type"]: m for m in movements}
    assert set(booked) == {"STORE_SHIPMENT", "UKR_SHIPMENT", "STORE_RETURN"}
    shipment = booked["STORE_SHIPMENT"]
    assert shipment["operation_date"] == TODAY.isoformat()
    assert shipment["lines"][0]["quantity"] == -75
    assert shipment["comment"] == (
        f"Transplaner {LAST_MONTH:%m.%Y} — według dni: "
        f"{day(1):%d.%m} -60, {day(2):%d.%m} -15"
    )
    assert (await stock_of(client, admin_headers))["EURO"] == -75 - 19 + 10


async def test_the_same_file_again_books_nothing(client, admin_headers, carriers):
    await book(client, admin_headers, FIRST)

    again = await book(client, admin_headers, FIRST)
    listed = await client.get("/movements", headers=admin_headers)

    assert again.status_code == 200
    assert {k["change"] for k in again.json()["summary"]["kinds"].values()} == {0}
    assert listed.headers["X-Total-Count"] == "3"


async def test_a_later_file_books_only_the_difference(client, admin_headers, carriers):
    await book(client, admin_headers, FIRST)
    later = build(
        # Day 2 of Białołęka was corrected down, day 3 is new.
        shipments={
            "002 BIALOLEKA": {1: 26, 2: 10, 3: 40},
            "003 GDANSK": {1: 34},
            UKR: {2: 19},
        },
        # A return confirmed later, dated by the store's sending day.
        returns={"002 BIALOLEKA": {1: 10, 2: 7}},
    )

    preview_body = (await preview(client, admin_headers, later)).json()
    await book(client, admin_headers, later)

    kinds = {k["type"]: k for k in preview_body["kinds"]}
    assert (kinds["STORE_SHIPMENT"]["booked"], kinds["STORE_SHIPMENT"]["change"]) == (
        75,
        -35,
    )
    assert [
        (d["day"], d["in_file"], d["booked"]) for d in kinds["STORE_SHIPMENT"]["days"]
    ] == [(day(2).isoformat(), 10, 15), (day(3).isoformat(), 40, 0)]
    assert kinds["STORE_RETURN"]["change"] == 7
    assert kinds["UKR_SHIPMENT"]["change"] == 0
    assert (await stock_of(client, admin_headers))["EURO"] == -75 - 19 + 10 - 35 + 7


async def test_only_the_chosen_operations_are_booked(client, admin_headers, carriers):
    await book(client, admin_headers, FIRST, "STORE_RETURN")

    first_types = [
        m["type"]
        for m in (await client.get("/movements", headers=admin_headers)).json()
    ]
    # The shipments were not booked, so a later upload still books them whole.
    later = (await preview(client, admin_headers, FIRST)).json()

    assert first_types == ["STORE_RETURN"]
    assert {k["type"]: k["change"] for k in later["kinds"]}["STORE_SHIPMENT"] == -75


async def test_nothing_chosen(client, admin_headers, carriers):
    response = await client.post(
        "/imports/transplaner",
        headers=admin_headers,
        files={"file": ("raport.xlsx", FIRST)},
        data={"type": []},
    )

    assert response.status_code in (400, 422)


async def test_an_operation_missing_from_the_file(client, admin_headers, carriers):
    response = await book(
        client, admin_headers, build(shipments={"A": {1: 1}}), "STORE_RETURN"
    )

    assert response.status_code == 400
    assert response.json() == {
        "message": "„Zwroty ze sklepów” nie występuje w tym pliku"
    }


async def test_a_carrier_without_a_type(client, admin_headers, carriers):
    response = await preview(
        client, admin_headers, build(shipments={"A": {1: 1}}, carrier="PAL CHEP")
    )

    assert response.status_code == 400
    assert response.json()["message"].startswith(
        "Nośnik „PAL CHEP” z pliku nie jest przypisany do typu nośnika"
    )


async def test_a_future_month_is_refused(client, admin_headers, carriers):
    next_month = (TODAY.replace(day=28) + timedelta(days=5)).replace(day=1)

    response = await preview(
        client,
        admin_headers,
        build(shipments={"A": {1: 1}}, month=next_month.strftime("%Y-%m")),
    )

    assert response.status_code == 400
    assert response.json() == {
        "message": f"Plik dotyczy miesiąca {next_month:%m.%Y}, który jeszcze się nie zaczął"
    }


async def test_the_list_of_imports(client, admin_headers, carriers):
    await book(client, admin_headers, FIRST)

    listed = (await client.get("/imports/transplaner", headers=admin_headers)).json()

    assert [(i["file_name"], i["user"]["username"]) for i in listed] == [
        ("raport.xlsx", "admin")
    ]


@pytest.mark.parametrize(
    ("role", "allowed"), [("operator", True), ("controller", False), ("viewer", False)]
)
async def test_who_may_import(client, session, carriers, role, allowed):
    await make_user(session, "jan", role=role)
    headers = await login(client, "jan")

    response = await book(client, headers, FIRST)

    assert response.status_code == (200 if allowed else 403)


async def test_the_name_in_transplaner_is_unique(client, admin_headers, carriers):
    response = await client.put(
        f"/carrier-types/{carriers.single}",
        headers=admin_headers,
        json={"transplaner_name": " PAL EUR/ EPAL "},
    )

    assert response.status_code == 409
    assert response.json() == {
        "message": "Nazwa „PAL EUR/ EPAL” z Transplanera jest już przypisana do „EURO”"
    }
