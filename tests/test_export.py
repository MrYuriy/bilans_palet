"""XLSX exports: the daily balance (D-1 for LM) and the movement register."""

from datetime import timedelta
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.services.day_close_service import DayCloseService
from tests.conftest import line, move, warehouse_today

TODAY = warehouse_today()
YESTERDAY = TODAY - timedelta(days=1)
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest.fixture
async def history(client, admin_headers, carriers):
    for kind, ln, doc in [
        ("PRODUCER_RECEIPT", line(carriers.euro, 100), "PZ/1"),
        ("ECOM_TRANSFER", line(carriers.euro, 30), "WZ/1"),
        ("OTHER_RECEIPT", line(carriers.single, 7), None),
    ]:
        response = await move(
            client, admin_headers, kind, ln, day=YESTERDAY, document_number=doc
        )
        assert response.status_code == 200, response.text


def sheet_of(response):
    assert response.status_code == 200, response.text
    assert response.headers["content-type"] == XLSX
    return load_workbook(BytesIO(response.content)).active


def rows_from(ws, first_row):
    return [
        list(row)
        for row in ws.iter_rows(min_row=first_row, values_only=True)
        if any(v is not None for v in row)
    ]


async def test_daily_balance_file(client, session, admin_headers, history):
    await DayCloseService(session).close_due_days()

    response = await client.get(
        "/balances/daily.xlsx",
        headers=admin_headers,
        params={"day": YESTERDAY.isoformat()},
    )

    assert (
        response.headers["content-disposition"]
        == f'attachment; filename="bilans_{YESTERDAY}.xlsx"'
    )
    ws = sheet_of(response)
    assert ws["A1"].value == "Bilans dzienny nośników"
    assert ws["B2"].value.date() == YESTERDAY
    assert ws["B3"].value.startswith("zamknięta")
    assert rows_from(ws, 8) == [
        [
            "Grupa",
            "Typ nośnika",
            "Stan otwarcia",
            "Przyjęcia",
            "Wydania",
            "Korekty",
            "Przeksięgowania",
            "Stan zamknięcia",
        ],
        ["EUR", "EURO", 0, 100, 30, 0, 0, 70],
        ["EUR", "RAZEM", 0, 100, 30, 0, 0, 70],
        ["Jednorazowe", "Jednorazowe", 0, 7, 0, 0, 0, 7],
        ["Jednorazowe", "RAZEM", 0, 7, 0, 0, 0, 7],
    ]


async def test_daily_balance_file_with_picked_types(
    client, admin_headers, carriers, history
):
    response = await client.get(
        "/balances/daily.xlsx",
        headers=admin_headers,
        params=[("day", YESTERDAY.isoformat()), ("type", carriers.damaged)],
    )

    assert [r[:2] for r in rows_from(sheet_of(response), 9)] == [
        ["EUR", "EURO uszkodzone"],
        ["EUR", "RAZEM"],
    ]


async def test_daily_balance_of_the_open_day_says_preview(
    client, admin_headers, history
):
    response = await client.get(
        "/balances/daily.xlsx", headers=admin_headers, params={"day": TODAY.isoformat()}
    )

    assert sheet_of(response)["B3"].value.startswith("podgląd")


async def test_movement_register_file(client, admin_headers, history):
    response = await client.get("/movements/export.xlsx", headers=admin_headers)

    assert response.headers["content-disposition"].startswith(
        'attachment; filename="rejestr_operacji_'
    )
    rows = rows_from(sheet_of(response), 1)
    assert rows[0] == [
        "Data operacji",
        "Operacja",
        "Typ nośnika",
        "Zmiana stanu",
        "Nr dokumentu",
        "Komentarz",
        "Użytkownik",
        "Wprowadzono",
    ]
    by_doc = {r[4]: r for r in rows[1:]}
    shipment = by_doc["WZ/1"]
    assert shipment[0].date() == YESTERDAY
    assert shipment[1:4] == ["Przesunięcie na ECOM", "EURO", -30]
    assert shipment[6] == "admin"
    assert len(rows) == 4


async def test_movement_register_file_follows_the_filters(
    client, admin_headers, history
):
    response = await client.get(
        "/movements/export.xlsx",
        headers=admin_headers,
        params={"type": "ECOM_TRANSFER"},
    )

    rows = rows_from(sheet_of(response), 2)
    assert [r[1] for r in rows] == ["Przesunięcie na ECOM"]


async def test_a_matching_count_is_one_row_with_zero(
    client, admin_headers, carriers, history
):
    await client.post(
        "/inventories",
        headers=admin_headers,
        json={"lines": [{"carrier_type_uuid": carriers.single, "counted_quantity": 7}]},
    )

    response = await client.get(
        "/movements/export.xlsx", headers=admin_headers, params={"type": "INVENTORY"}
    )

    rows = rows_from(sheet_of(response), 2)
    assert [r[1:4] for r in rows] == [["Inwentaryzacja", "—", 0]]


async def test_exports_need_a_login(client):
    for url in ("/balances/daily.xlsx", "/movements/export.xlsx"):
        assert (await client.get(url)).status_code == 401
